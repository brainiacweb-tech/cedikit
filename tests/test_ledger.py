import json
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest

from cedikit import sms
from cedikit.ledger import Ledger, LedgerEntry, Rule, expected_balance

D = Decimal
Samples = dict[str, dict[str, Any]]


def _msgs(genuine: Samples, *ids: str) -> list[tuple[str, str]]:
    return [(genuine[i]["text"], genuine[i]["sender"]) for i in ids]


@pytest.fixture
def telecel_day(genuine: Samples) -> Ledger:
    """Deposit 80.00 -> airtime 30.00 -> airtime notice (not counted)."""
    return Ledger.from_messages(
        _msgs(
            genuine,
            "telecel_cash_in",
            "telecel_airtime_purchase",
            "telecel_airtime_received_notice",
        )
    )


@pytest.fixture
def everything(genuine: Samples) -> Ledger:
    return Ledger.from_messages([(s["text"], s["sender"]) for s in genuine.values()])


def test_notices_and_unrecognised_are_kept_apart(genuine: Samples) -> None:
    ledger = Ledger.from_messages(
        [*_msgs(genuine, "telecel_cash_in", "telecel_airtime_received_notice"), "hello"]
    )
    assert len(ledger) == 1
    assert len(ledger.notices) == 1
    assert [r.raw for r in ledger.unrecognised] == ["hello"]


def test_duplicate_transaction_ids_are_dropped(genuine: Samples) -> None:
    ledger = Ledger.from_messages(_msgs(genuine, "mtn_cash_out", "mtn_cash_out"))
    assert len(ledger) == 1
    assert ledger.duplicates == 1


def test_input_forms(genuine: Samples) -> None:
    text = genuine["mtn_cash_out"]["text"]
    when = datetime(2026, 1, 20, 10, 4, tzinfo=sms.GHANA_TIME)
    parsed = sms.parse(genuine["mtn_received"]["text"], "MobileMoney")
    assert parsed.transaction is not None
    ledger = Ledger.from_messages(
        [
            {"text": text, "received_at": when},
            (genuine["mtn_cash_in_same_wallet"]["text"], None),
            parsed,
            sms.parse(genuine["mtn_loan_repayment"]["text"]).transaction,  # type: ignore[list-item]
            genuine["mtn_payment_dotted_payee"]["text"],
        ],
        sender="MobileMoney",
    )
    assert len(ledger) == 5
    assert ledger.transactions[0].timestamp == when
    assert all(t.sender in ("MobileMoney", None) for t in ledger)
    with pytest.raises(TypeError):
        Ledger.from_messages([42])  # type: ignore[list-item]


def test_sorted_by_time_only_when_all_dated(genuine: Samples) -> None:
    dated = Ledger.from_messages(
        _msgs(genuine, "telecel_received_same_network", "telecel_sent_same_network")
    )
    assert [t.type.value for t in dated] == ["SENT", "RECEIVED"]
    mixed = Ledger.from_messages(_msgs(genuine, "telecel_received_same_network", "mtn_cash_out"))
    assert [t.network for t in mixed] == ["TELECEL", "MTN"]


def test_summary(telecel_day: Ledger) -> None:
    s = telecel_day.summary()
    assert (s.transactions, s.total_in, s.total_out, s.net) == (
        2,
        D("80.00"),
        D("30.00"),
        D("50.00"),
    )
    assert s.closing_balances == {"TELECEL": D("50.37")}
    assert s.first is not None and s.first.day == 8
    text = str(s)
    assert "08 Feb 2026" in text and "GH₵ 80.00" in text and "Last TELECEL balance" in text


def test_summary_without_dates(genuine: Samples) -> None:
    s = Ledger.from_messages(_msgs(genuine, "mtn_cash_out")).summary()
    assert s.undated == 1 and s.first is None
    assert "dates unknown" in str(s) and "have no date" in str(s)
    assert Ledger().summary().transactions == 0


def test_fees_and_taxes(genuine: Samples) -> None:
    s = Ledger.from_messages(_msgs(genuine, "mtn_sent_to_telecel")).summary()
    assert (s.fees, s.taxes, s.net) == (D("10.00"), D("15.00"), D("-1525.00"))


def test_cash_flow(everything: Ledger) -> None:
    months = everything.cash_flow("month")
    assert [p.start.isoformat() for p in months][:2] == ["2024-03-01", "2026-01-01"]
    jan = months[1]
    assert jan.inflow > 0 and jan.net == jan.inflow - jan.outflow - jan.fees - jan.taxes
    weeks = everything.cash_flow("week")
    assert all(p.start.weekday() == 0 for p in weeks)
    days = everything.cash_flow("day")
    assert sum(p.count for p in days) == len(everything) - everything.summary().undated
    with pytest.raises(ValueError, match="Unknown period"):
        everything.cash_flow("year")  # type: ignore[arg-type]


def test_top_counterparties(genuine: Samples) -> None:
    ledger = Ledger.from_messages(
        _msgs(
            genuine,
            "mtn_cash_out",
            "mtn_cash_in_same_wallet",
            "mtn_received",
            "mtn_sent_to_telecel",
        )
    )
    by_count = ledger.top_counterparties(2)
    assert by_count[0].name == "ADOM ELECTRICALS" and by_count[0].count == 2
    assert by_count[0].total_in == D("300.00") and by_count[0].total_out == D("40.00")
    by_value = ledger.top_counterparties(1, by="value")
    assert by_value[0].name == "YAW DARKO" and by_value[0].total == D("1500.00")
    with pytest.raises(ValueError):
        ledger.top_counterparties(by="size")  # type: ignore[arg-type]


def test_categorise_defaults(everything: Ledger) -> None:
    categorised = everything.categorise()
    by_template = {e.transaction.template: e.category for e in categorised.entries}
    assert by_template["mtn_payment"] in ("loan repayment", "payments & supplies")
    assert by_template["mtn_cash_out"] == "cash withdrawal"
    assert by_template["telecel_airtime_purchase"] == "airtime & data"
    loan = next(e for e in categorised.entries if "Loan" in (e.transaction.counterparty.name or ""))  # type: ignore[union-attr]
    assert loan.category == "loan repayment"
    totals = categorised.category_totals()
    assert list(totals.values()) == sorted(totals.values(), reverse=True)
    assert categorised.notices == everything.notices


def test_categorise_custom_rules(everything: Ledger) -> None:
    rules = [
        {"category": "big transfers", "types": ["SENT"], "min_amount": "1000"},
        {"category": "rent", "reference_contains": "rent"},
        Rule("adom", name_contains="adom", max_amount=D("100")),
        {"category": "esi", "phone": "0509876543"},
    ]
    cats = {e.transaction.transaction_id: e.category for e in everything.categorise(rules).entries}
    assert cats["40123456789"] == "big transfers"
    assert cats["0000019301122334"] == "rent"  # reference "Rent" beats phone rule order
    assert cats["10493827561"] == "adom"  # cash out 40.00 to ADOM
    assert cats["10502294418"] == "cash deposit"  # ADOM 300.00 is above max_amount
    bare = everything.categorise(include_defaults=False)
    assert {e.category for e in bare.entries} == {"uncategorised"}
    assert Ledger(everything.entries).category_totals().keys() == {"uncategorised"}


def test_balance_gaps(genuine: Samples) -> None:
    chain = Ledger.from_messages(
        _msgs(
            genuine, "mtn_cash_out", "mtn_cash_in_same_wallet", "mtn_received", "mtn_loan_repayment"
        )
    )
    gaps = chain.balance_gaps()
    # cash_out -> cash_in adds up; cash_in -> received does not (different wallets in fixtures)
    assert len(gaps) == 1
    gap = gaps[0]
    assert gap.before.template == "mtn_cash_in" and gap.after.template == "mtn_received"
    assert gap.expected == D("358.10") and gap.actual == D("36.00")
    assert gap.difference == D("-322.10")


def test_expected_balance() -> None:
    assert expected_balance(D("10"), D("5"), "in") == D("15")
    assert expected_balance(D("10"), D("5"), "out", D("0.10"), D("0.05")) == D("4.85")


def test_csv_round_trip(everything: Ledger, tmp_path: Path) -> None:
    categorised = everything.categorise()
    path = categorised.export(tmp_path / "ledger.csv")
    loaded = Ledger.from_csv(path)
    assert loaded.to_rows() == categorised.to_rows()


def test_json_export(telecel_day: Ledger, tmp_path: Path) -> None:
    data = json.loads(telecel_day.export(tmp_path / "l.json").read_text("utf-8"))
    assert data["summary"]["total_in"] == "80.00"
    assert len(data["transactions"]) == 2


def test_excel_export(everything: Ledger, tmp_path: Path) -> None:
    import openpyxl

    path = everything.categorise().export(tmp_path / "ledger.xlsx")
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ["Transactions", "Summary", "Cash flow", "Categories"]
    tx = book["Transactions"]
    assert tx.max_row == len(everything) + 1
    assert tx["A1"].value == "timestamp"
    plain = everything.export(tmp_path / "plain.xlsx")
    assert "Categories" not in openpyxl.load_workbook(plain).sheetnames


def test_export_unknown_format(telecel_day: Ledger, tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="Unknown export format"):
        telecel_day.export(tmp_path / "ledger.pdf")


def test_dataframe(everything: Ledger) -> None:
    frame = everything.to_dataframe()
    assert len(frame) == len(everything)
    assert isinstance(frame["amount"].iloc[0], Decimal)
    assert str(frame["timestamp"].dtype).startswith("datetime64")


def test_plot(everything: Ledger, tmp_path: Path) -> None:
    fig = everything.plot(path=tmp_path / "flow.png")
    assert (tmp_path / "flow.png").stat().st_size > 1000
    assert fig.axes[0].get_title() == "Cash flow per month"
    everything.categorise().plot("categories", path=tmp_path / "cats.png")
    with pytest.raises(ValueError):
        everything.plot("pie")  # type: ignore[arg-type]


def test_entries_can_be_passed_directly(telecel_day: Ledger) -> None:
    copy = Ledger([LedgerEntry(t, "x") for t in telecel_day])
    assert [e.category for e in copy.entries] == ["x", "x"]
