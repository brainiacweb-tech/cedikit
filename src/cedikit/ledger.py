"""Turn Mobile Money transactions into a ledger with summaries, cash flow and exports.

Example:
    >>> from cedikit.ledger import Ledger
    >>> ledger = Ledger.from_messages([
    ...     ("Cash In received for GHS 100.00 from ADOM VENTURES . Current Balance GHS 120.00 "
    ...      "Available Balance GHS 120.00. Transaction ID: 10000000001. Fee charged: GHS 0.",
    ...      "MobileMoney"),
    ... ])
    >>> ledger.summary().total_in
    Decimal('100.00')
"""

from __future__ import annotations

import csv
import json
from collections.abc import Iterable, Iterator, Mapping
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING, Any, Literal

from cedikit import money, phone
from cedikit.sms.models import Counterparty, ParseResult, Transaction, TransactionType
from cedikit.sms.parser import Parser, _default_parser

if TYPE_CHECKING:
    import pandas as pd
    from matplotlib.figure import Figure

__all__ = [
    "DEFAULT_RULES",
    "BalanceGap",
    "CashFlowPeriod",
    "CounterpartyTotal",
    "Ledger",
    "LedgerEntry",
    "LedgerSummary",
    "Rule",
    "expected_balance",
    "read_messages",
    "split_messages",
]

Period = Literal["day", "week", "month"]
MessageInput = str | tuple[str, str | None] | Mapping[str, Any] | ParseResult | Transaction
ZERO = Decimal("0.00")

COLUMNS = [
    "timestamp",
    "network",
    "type",
    "direction",
    "amount",
    "fee",
    "tax",
    "balance",
    "counterparty_name",
    "counterparty_phone",
    "counterparty_network",
    "reference",
    "transaction_id",
    "category",
    "template",
    "confidence",
    "sender",
    "raw",
]


def _require(module: str, extra: str) -> Any:
    try:
        return __import__(module)
    except ImportError:  # pragma: no cover - depends on the environment
        raise ImportError(
            f"This feature needs {module}. Install it with: pip install 'cedikit[{extra}]'"
        ) from None


def expected_balance(
    previous: Decimal,
    amount: Decimal,
    direction: Literal["in", "out"],
    fee: Decimal | None = None,
    tax: Decimal | None = None,
) -> Decimal:
    """The wallet balance after a transaction, given the balance before it.

    Fees and taxes always leave the wallet, whichever way the amount moves.

    Example:
        >>> expected_balance(Decimal("100.00"), Decimal("40.00"), "out", Decimal("0.20"))
        Decimal('59.80')
    """
    signed = amount if direction == "in" else -amount
    return previous + signed - (fee or ZERO) - (tax or ZERO)


@dataclass(frozen=True)
class Rule:
    """A categorisation rule. Every condition given must match.

    Example:
        >>> Rule("stock", name_contains="wholesale").category
        'stock'
    """

    category: str
    types: frozenset[TransactionType] | None = None
    name_contains: str | None = None
    reference_contains: str | None = None
    phone: str | None = None
    min_amount: Decimal | None = None
    max_amount: Decimal | None = None

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> Rule:
        """Build a rule from plain data, e.g. one entry of a YAML rules file."""
        types = data.get("types")
        return cls(
            category=str(data["category"]),
            types=frozenset(TransactionType(t) for t in types) if types else None,
            name_contains=data.get("name_contains"),
            reference_contains=data.get("reference_contains"),
            phone=phone.normalise(data["phone"]) if data.get("phone") else None,
            min_amount=money.parse(data["min_amount"]) if "min_amount" in data else None,
            max_amount=money.parse(data["max_amount"]) if "max_amount" in data else None,
        )

    def matches(self, tx: Transaction) -> bool:
        cp = tx.counterparty or Counterparty(None)
        checks = [
            self.types is None or tx.type in self.types,
            self.name_contains is None
            or self.name_contains.casefold() in (cp.name or "").casefold(),
            self.reference_contains is None
            or self.reference_contains.casefold() in (tx.reference or "").casefold(),
            self.phone is None or self.phone == cp.phone,
            self.min_amount is None or tx.amount >= self.min_amount,
            self.max_amount is None or tx.amount <= self.max_amount,
        ]
        return all(checks)


T = TransactionType
DEFAULT_RULES: tuple[Rule, ...] = (
    Rule("loan repayment", name_contains="loan"),
    Rule("airtime & data", types=frozenset({T.AIRTIME})),
    Rule("bills", types=frozenset({T.BILL})),
    Rule("reversals", types=frozenset({T.REVERSAL})),
    Rule("cash deposit", types=frozenset({T.CASH_IN})),
    Rule("cash withdrawal", types=frozenset({T.CASH_OUT})),
    Rule("payments & supplies", types=frozenset({T.MERCHANT})),
    Rule("sales & money received", types=frozenset({T.RECEIVED})),
    Rule("money sent", types=frozenset({T.SENT})),
)
"""Applied after your own rules. The first matching rule wins."""


@dataclass(frozen=True)
class LedgerEntry:
    transaction: Transaction
    category: str | None = None


@dataclass(frozen=True)
class LedgerSummary:
    """Totals for a ledger. ``net`` is money in minus money out, fees and taxes."""

    transactions: int
    total_in: Decimal
    total_out: Decimal
    fees: Decimal
    taxes: Decimal
    net: Decimal
    first: datetime | None
    last: datetime | None
    closing_balances: dict[str, Decimal]
    undated: int

    def __str__(self) -> str:
        span = (
            f"{self.first:%d %b %Y} to {self.last:%d %b %Y}"
            if self.first and self.last
            else "dates unknown"
        )
        lines = [
            f"{self.transactions} transactions ({span})",
            f"  Money in:  {money.format(self.total_in)}",
            f"  Money out: {money.format(self.total_out)}",
            f"  Fees:      {money.format(self.fees)}",
            f"  Taxes:     {money.format(self.taxes)}",
            f"  Net:       {money.format(self.net)}",
        ]
        for network, balance in self.closing_balances.items():
            lines.append(f"  Last {network} balance: {money.format(balance)}")
        if self.undated:
            lines.append(f"  ({self.undated} transactions have no date)")
        return "\n".join(lines)


@dataclass(frozen=True)
class CashFlowPeriod:
    start: date
    inflow: Decimal
    outflow: Decimal
    fees: Decimal
    taxes: Decimal
    count: int

    @property
    def net(self) -> Decimal:
        return self.inflow - self.outflow - self.fees - self.taxes


@dataclass(frozen=True)
class CounterpartyTotal:
    name: str | None
    phone: str | None
    count: int
    total_in: Decimal
    total_out: Decimal

    @property
    def total(self) -> Decimal:
        return self.total_in + self.total_out


@dataclass(frozen=True)
class BalanceGap:
    """Two consecutive transactions whose balances don't add up.

    Usually a message is missing between them (e.g. an automatic loan
    deduction that wasn't imported); occasionally one of them is fake.
    """

    before: Transaction
    after: Transaction
    expected: Decimal
    actual: Decimal

    @property
    def difference(self) -> Decimal:
        return self.actual - self.expected


def _period_start(moment: datetime, period: Period) -> date:
    day = moment.date()
    if period == "day":
        return day
    if period == "week":
        return day - timedelta(days=day.weekday())
    return day.replace(day=1)


class Ledger:
    """An ordered collection of wallet transactions.

    Only transactions that move wallet money are included: notices such as
    "you have received airtime" are kept in :attr:`notices`, unrecognised
    messages in :attr:`unrecognised`, and repeated transaction IDs are dropped
    (counted in :attr:`duplicates`).

    Transactions are sorted by time when all of them have a timestamp;
    otherwise they keep the order you gave them in (phones list SMS oldest
    first, so pass them that way).
    """

    def __init__(
        self,
        transactions: Iterable[Transaction | LedgerEntry] = (),
        *,
        unrecognised: Iterable[ParseResult] = (),
    ) -> None:
        self.entries: list[LedgerEntry] = []
        self.notices: list[Transaction] = []
        self.unrecognised: list[ParseResult] = list(unrecognised)
        self.duplicates = 0
        seen: set[tuple[str, str]] = set()
        for item in transactions:
            entry = item if isinstance(item, LedgerEntry) else LedgerEntry(item)
            tx = entry.transaction
            if not tx.affects_wallet:
                self.notices.append(tx)
                continue
            if tx.transaction_id:
                key = (tx.network, tx.transaction_id)
                if key in seen:
                    self.duplicates += 1
                    continue
                seen.add(key)
            self.entries.append(entry)
        if self.entries and all(e.transaction.timestamp for e in self.entries):
            self.entries.sort(key=lambda e: e.transaction.timestamp)  # type: ignore[arg-type, return-value]

    # -- building ---------------------------------------------------------

    @classmethod
    def from_messages(
        cls,
        messages: Iterable[MessageInput],
        sender: str | None = None,
        *,
        parser: Parser | None = None,
    ) -> Ledger:
        """Parse SMS messages into a ledger.

        Each message can be the text alone, a ``(text, sender)`` tuple, or a
        mapping with ``text`` and optional ``sender`` and ``received_at`` keys.
        ``sender`` applies to messages that don't name their own.
        """
        parser = parser or _default_parser()
        transactions: list[Transaction] = []
        unrecognised: list[ParseResult] = []
        for item in messages:
            if isinstance(item, Transaction):
                transactions.append(item)
                continue
            if isinstance(item, ParseResult):
                result = item
            else:
                text, msg_sender, received_at = _unpack(item, sender)
                result = parser.parse(text, msg_sender, received_at)
            if result.transaction is None:
                unrecognised.append(result)
            else:
                transactions.append(result.transaction)
        return cls(transactions, unrecognised=unrecognised)

    @classmethod
    def from_csv(cls, path: str | Path) -> Ledger:
        """Load a ledger saved with :meth:`export` as CSV."""
        with Path(path).open(newline="", encoding="utf-8") as fh:
            rows = list(csv.DictReader(fh))
        return cls(_entry_from_row(row) for row in rows)

    # -- basics -------------------------------------------------------------

    def __len__(self) -> int:
        return len(self.entries)

    def __iter__(self) -> Iterator[Transaction]:
        return (e.transaction for e in self.entries)

    @property
    def transactions(self) -> list[Transaction]:
        return [e.transaction for e in self.entries]

    # -- analysis -------------------------------------------------------------

    def summary(self) -> LedgerSummary:
        """Money in, money out, fees, taxes and net flow."""
        txs = self.transactions
        total_in = sum((t.amount for t in txs if t.type.direction == "in"), ZERO)
        total_out = sum((t.amount for t in txs if t.type.direction == "out"), ZERO)
        fees = sum((t.fee or ZERO for t in txs), ZERO)
        taxes = sum((t.tax or ZERO for t in txs), ZERO)
        dated = [t.timestamp for t in txs if t.timestamp]
        closing: dict[str, Decimal] = {}
        for t in txs:
            if t.balance is not None:
                closing[t.network] = t.balance
        return LedgerSummary(
            transactions=len(txs),
            total_in=total_in,
            total_out=total_out,
            fees=fees,
            taxes=taxes,
            net=total_in - total_out - fees - taxes,
            first=min(dated) if dated else None,
            last=max(dated) if dated else None,
            closing_balances=closing,
            undated=len(txs) - len(dated),
        )

    def cash_flow(self, period: Period = "month") -> list[CashFlowPeriod]:
        """Money in and out per day, week (from Monday) or month.

        Transactions without a timestamp are left out; see ``summary().undated``.
        """
        if period not in ("day", "week", "month"):
            raise ValueError(f"Unknown period {period!r}; use 'day', 'week' or 'month'")
        buckets: dict[date, list[Transaction]] = {}
        for t in self:
            if t.timestamp:
                buckets.setdefault(_period_start(t.timestamp, period), []).append(t)
        return [
            CashFlowPeriod(
                start=start,
                inflow=sum((t.amount for t in txs if t.type.direction == "in"), ZERO),
                outflow=sum((t.amount for t in txs if t.type.direction == "out"), ZERO),
                fees=sum((t.fee or ZERO for t in txs), ZERO),
                taxes=sum((t.tax or ZERO for t in txs), ZERO),
                count=len(txs),
            )
            for start, txs in sorted(buckets.items())
        ]

    def top_counterparties(
        self, n: int = 5, by: Literal["count", "value"] = "count"
    ) -> list[CounterpartyTotal]:
        """The people and businesses you deal with most, by number or value."""
        groups: dict[str, list[Transaction]] = {}
        for t in self:
            cp = t.counterparty
            if cp is None:
                continue
            key = cp.phone or (cp.name or "").upper()
            groups.setdefault(key, []).append(t)
        totals = [
            CounterpartyTotal(
                name=next((t.counterparty.name for t in txs if t.counterparty and t.counterparty.name), None),  # noqa: E501
                phone=txs[0].counterparty.phone if txs[0].counterparty else None,
                count=len(txs),
                total_in=sum((t.amount for t in txs if t.type.direction == "in"), ZERO),
                total_out=sum((t.amount for t in txs if t.type.direction == "out"), ZERO),
            )
            for txs in groups.values()
        ]  # fmt: skip
        if by == "count":
            totals.sort(key=lambda c: (c.count, c.total), reverse=True)
        elif by == "value":
            totals.sort(key=lambda c: (c.total, c.count), reverse=True)
        else:
            raise ValueError("by must be 'count' or 'value'")
        return totals[:n]

    def categorise(
        self,
        rules: Iterable[Rule | Mapping[str, Any]] = (),
        *,
        include_defaults: bool = True,
    ) -> Ledger:
        """Return a copy with a category on every transaction.

        Your ``rules`` are tried first, then :data:`DEFAULT_RULES`. The first
        match wins; anything unmatched is ``"uncategorised"``.

        Example:
            >>> rules = [{"category": "stock", "name_contains": "wholesale"}]
            >>> categorised = ledger.categorise(rules)  # doctest: +SKIP
        """
        compiled = [r if isinstance(r, Rule) else Rule.from_dict(r) for r in rules]
        if include_defaults:
            compiled.extend(DEFAULT_RULES)
        entries = [
            LedgerEntry(
                e.transaction,
                next((r.category for r in compiled if r.matches(e.transaction)), "uncategorised"),
            )
            for e in self.entries
        ]
        new = Ledger(entries, unrecognised=self.unrecognised)
        new.notices, new.duplicates = list(self.notices), self.duplicates
        return new

    def category_totals(self) -> dict[str, Decimal]:
        """Total amount per category, largest first. Call :meth:`categorise` first."""
        totals: dict[str, Decimal] = {}
        for e in self.entries:
            key = e.category or "uncategorised"
            totals[key] = totals.get(key, ZERO) + e.transaction.amount
        return dict(sorted(totals.items(), key=lambda kv: kv[1], reverse=True))

    def balance_gaps(self, tolerance: Decimal = Decimal("0.01")) -> list[BalanceGap]:
        """Consecutive transactions (per network) whose balances don't add up."""
        gaps: list[BalanceGap] = []
        previous: dict[str, Transaction] = {}
        for t in self:
            if t.balance is None:
                continue
            before = previous.get(t.network)
            if before is not None and before.balance is not None:
                expected = expected_balance(
                    before.balance, t.amount, t.type.direction, t.fee, t.tax
                )
                if abs(expected - t.balance) > tolerance:
                    gaps.append(BalanceGap(before, t, expected, t.balance))
            previous[t.network] = t
        return gaps

    # -- export ---------------------------------------------------------------

    def to_rows(self) -> list[dict[str, str]]:
        """One flat dict of strings per transaction (the CSV export format)."""
        return [_row(e) for e in self.entries]

    def to_dataframe(self) -> pd.DataFrame:
        """A pandas DataFrame with Decimal money columns and parsed timestamps."""
        pandas = _require("pandas", "pandas")
        frame: pd.DataFrame = pandas.DataFrame(self.to_rows(), columns=COLUMNS)
        for col in ("amount", "fee", "tax", "balance"):
            frame[col] = frame[col].map(lambda v: Decimal(v) if v else None)
        frame["timestamp"] = pandas.to_datetime(frame["timestamp"].replace("", None), utc=True)
        return frame

    def export(
        self, path: str | Path, format: Literal["csv", "xlsx", "json"] | None = None
    ) -> Path:
        """Save as CSV, Excel (``.xlsx``) or JSON; the format follows the extension.

        Excel files get Transactions, Summary, Cash flow and (if categorised)
        Categories sheets.
        """
        path = Path(path)
        fmt = format or path.suffix.lstrip(".").lower()
        if fmt == "csv":
            with path.open("w", newline="", encoding="utf-8") as fh:
                writer = csv.DictWriter(fh, fieldnames=COLUMNS)
                writer.writeheader()
                writer.writerows(self.to_rows())
        elif fmt == "json":
            s = self.summary()
            payload = {
                "summary": {
                    "transactions": s.transactions,
                    "total_in": str(s.total_in),
                    "total_out": str(s.total_out),
                    "fees": str(s.fees),
                    "taxes": str(s.taxes),
                    "net": str(s.net),
                    "closing_balances": {k: str(v) for k, v in s.closing_balances.items()},
                },
                "transactions": self.to_rows(),
            }
            path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        elif fmt == "xlsx":
            self._export_xlsx(path)
        else:
            raise ValueError(f"Unknown export format {fmt!r}; use csv, xlsx or json")
        return path

    def _export_xlsx(self, path: Path) -> None:
        openpyxl = _require("openpyxl", "excel")
        from openpyxl.styles import Font

        money_format = "#,##0.00"
        bold = Font(bold=True)
        book = openpyxl.Workbook()

        sheet = book.active
        sheet.title = "Transactions"
        headers = [c for c in COLUMNS if c != "raw"]
        sheet.append(headers)
        for e in self.entries:
            t = e.transaction
            cp = t.counterparty or Counterparty(None)
            sheet.append([
                t.timestamp.replace(tzinfo=None) if t.timestamp else None,
                t.network, t.type.value, t.type.direction, t.amount, t.fee, t.tax,
                t.balance, cp.name, cp.phone, cp.network, t.reference, t.transaction_id,
                e.category, t.template, t.confidence, t.sender,
            ])  # fmt: skip
        for row in sheet.iter_rows(min_row=2):
            row[0].number_format = "yyyy-mm-dd hh:mm"
            for cell in row[4:8]:
                cell.number_format = money_format

        s = self.summary()
        summary = book.create_sheet("Summary")
        for label, value in [
            ("Transactions", s.transactions),
            ("Money in", s.total_in),
            ("Money out", s.total_out),
            ("Fees", s.fees),
            ("Taxes", s.taxes),
            ("Net", s.net),
            *((f"Last {net} balance", bal) for net, bal in s.closing_balances.items()),
        ]:
            summary.append([label, value])
            if isinstance(value, Decimal):
                summary.cell(summary.max_row, 2).number_format = money_format

        flow = book.create_sheet("Cash flow")
        flow.append(["Month", "In", "Out", "Fees", "Taxes", "Net", "Transactions"])
        for p in self.cash_flow("month"):
            flow.append([p.start, p.inflow, p.outflow, p.fees, p.taxes, p.net, p.count])
            flow.cell(flow.max_row, 1).number_format = "mmm yyyy"
            for col in range(2, 7):
                flow.cell(flow.max_row, col).number_format = money_format

        if any(e.category for e in self.entries):
            cats = book.create_sheet("Categories")
            cats.append(["Category", "Total"])
            for name, total in self.category_totals().items():
                cats.append([name, total])
                cats.cell(cats.max_row, 2).number_format = money_format

        for ws in book.worksheets:
            for cell in ws[1]:
                cell.font = bold
            ws.freeze_panes = "A2"
            for column in ws.columns:
                width = max(len(str(c.value or "")) for c in column[:200])
                ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 45)
        book.save(path)

    def plot(
        self,
        kind: Literal["cash_flow", "categories"] = "cash_flow",
        period: Period = "month",
        path: str | Path | None = None,
    ) -> Figure:
        """Draw a quick chart; saved to ``path`` if given. Needs matplotlib."""
        _require("matplotlib", "charts")
        from matplotlib.figure import Figure

        fig = Figure(figsize=(8, 4.5), layout="constrained")
        ax = fig.add_subplot()
        if kind == "cash_flow":
            flows = self.cash_flow(period)
            labels = [p.start.isoformat() for p in flows]
            xs = range(len(flows))
            ax.bar([x - 0.2 for x in xs], [float(p.inflow) for p in flows], 0.4, label="In")
            ax.bar([x + 0.2 for x in xs], [float(p.outflow) for p in flows], 0.4, label="Out")
            ax.set_xticks(list(xs), labels, rotation=45, ha="right")
            ax.set_title(f"Cash flow per {period}")
            ax.legend()
        elif kind == "categories":
            totals = self.category_totals()
            ax.barh(list(totals)[::-1], [float(v) for v in totals.values()][::-1])
            ax.set_title("Amount per category")
        else:
            raise ValueError("kind must be 'cash_flow' or 'categories'")
        ax.set_ylabel("GHS" if kind == "cash_flow" else "")
        if path is not None:
            fig.savefig(path, dpi=150)
        return fig


def split_messages(text: str) -> list[str]:
    """Split pasted text into messages: one message per paragraph (blank line between).

    Example:
        >>> split_messages("first message\\n\\n  second\\nmessage  \\n\\n\\n")
        ['first message', 'second\\nmessage']
    """
    blocks = text.replace("\r\n", "\n").split("\n\n")
    return [b.strip() for b in blocks if b.strip()]


def read_messages(path: str | Path) -> list[dict[str, str]]:
    """Read messages from a file, ready for :meth:`Ledger.from_messages`.

    A ``.csv`` file needs a ``text`` column, and may have ``sender`` and
    ``received_at`` (ISO date-time) columns. Any other file is plain text with
    one message per paragraph.
    """
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as fh:
            return [row for row in csv.DictReader(fh) if row.get("text")]
    return [{"text": m} for m in split_messages(path.read_text(encoding="utf-8-sig"))]


def _unpack(item: MessageInput, default_sender: str | None) -> tuple[str, str | None, Any]:
    if isinstance(item, str):
        return item, default_sender, None
    if isinstance(item, tuple):
        return item[0], item[1] if item[1] is not None else default_sender, None
    if isinstance(item, Mapping):
        received_at = item.get("received_at") or None
        if isinstance(received_at, str):
            received_at = datetime.fromisoformat(received_at)
        return item["text"], item.get("sender") or default_sender, received_at
    raise TypeError(f"Unsupported message type {type(item).__name__}")


def _text(value: object) -> str:
    return "" if value is None else str(value)


def _row(entry: LedgerEntry) -> dict[str, str]:
    t = entry.transaction
    cp = t.counterparty or Counterparty(None)
    return {
        "timestamp": t.timestamp.isoformat() if t.timestamp else "",
        "network": t.network,
        "type": t.type.value,
        "direction": t.type.direction,
        "amount": str(t.amount),
        "fee": _text(t.fee),
        "tax": _text(t.tax),
        "balance": _text(t.balance),
        "counterparty_name": _text(cp.name),
        "counterparty_phone": _text(cp.phone),
        "counterparty_network": _text(cp.network),
        "reference": _text(t.reference),
        "transaction_id": _text(t.transaction_id),
        "category": _text(entry.category),
        "template": t.template,
        "confidence": str(t.confidence),
        "sender": _text(t.sender),
        "raw": t.raw,
    }


def _entry_from_row(row: Mapping[str, str]) -> LedgerEntry:
    def dec(key: str) -> Decimal | None:
        return Decimal(row[key]) if row.get(key) else None

    def opt(key: str) -> str | None:
        return row.get(key) or None

    cp = Counterparty(
        opt("counterparty_name"), opt("counterparty_phone"), opt("counterparty_network")
    )
    tx = Transaction(
        network=row["network"],
        type=TransactionType(row["type"]),
        amount=Decimal(row["amount"]),
        fee=dec("fee"),
        tax=dec("tax"),
        counterparty=cp if (cp.name or cp.phone) else None,
        transaction_id=opt("transaction_id"),
        reference=opt("reference"),
        balance=dec("balance"),
        available_balance=None,
        timestamp=datetime.fromisoformat(row["timestamp"]) if row.get("timestamp") else None,
        confidence=float(row.get("confidence") or 1.0),
        template=row.get("template") or "csv",
        sender=opt("sender"),
        raw=row.get("raw", ""),
    )
    return LedgerEntry(tx, opt("category"))
